import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from functools import wraps
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth import login as auth_login, logout as auth_logout
from django.http import HttpResponse
from django.db.models import Q, Value
from django.db.models.functions import Lower, Replace
from django.core.paginator import Paginator
from .models import Tomo, Serie, Autor, Editorial, Demografia
from .forms import (
    TomoForm, SerieForm, AutorForm, EditorialForm, DemografiaForm,
    RegistroUsuarioForm, LoginForm
)

# Decorador para limitar acciones solo a usuarios con rol Administrador
def solo_admin(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            messages.warning(request, 'Debes iniciar sesión para realizar esta acción.')
            return redirect(f"/login/?next={request.path}")
        if not request.user.is_staff and not request.user.is_superuser:
            messages.error(request, 'Acceso denegado: Se requieren permisos de Administrador para realizar esta acción.')
            return redirect('inicio')
        return view_func(request, *args, **kwargs)
    return _wrapped_view

# 1. Página inicial (muestra resumen y destacados)
def pagina_inicio(request):
    total_tomos = Tomo.objects.count()
    total_series = Serie.objects.count()
    ultimos_tomos = Tomo.objects.select_related('serie', 'editorial').order_by('-id')[:4]
    
    return render(request, 'inicio.html', {
        'total_tomos': total_tomos,
        'total_series': total_series,
        'ultimos_tomos': ultimos_tomos,
    })

# Función auxiliar para filtrar tomos según parámetros GET (Requisito 11)
def obtener_tomos_filtrados(request):
    """
    Retorna un queryset de Tomo filtrado según los parámetros GET recibidos.
    Se reutiliza en listar_tomos, exportar_tomos_excel y exportar_tomos_csv.
    """
    queryset = Tomo.objects.select_related('serie', 'editorial', 'serie__demografia', 'serie__autor').order_by('-id')

    # Búsqueda por texto (Título de serie, ISBN, Autor o Editorial)
    q = request.GET.get('q', '').strip()
    if q:
        # Normalizar columnas en la consulta (sin espacios, sin guiones, minúsculas)
        queryset = queryset.annotate(
            titulo_sin_espacios=Replace(Lower('serie__titulo'), Value(' '), Value('')),
            autor_sin_espacios=Replace(Lower('serie__autor__nombre'), Value(' '), Value('')),
            editorial_sin_espacios=Replace(Lower('editorial__nombre'), Value(' '), Value('')),
            isbn_sin_guiones=Replace(Replace('isbn', Value('-'), Value('')), Value(' '), Value(''))
        )

        q_str = q.lower()
        q_sin_espacios = q_str.replace(' ', '').replace('-', '')

        # Filtro flexible: coincide con o sin espacios, mayúsculas o minúsculas
        filtro_texto = (
            Q(serie__titulo__icontains=q) |
            Q(titulo_sin_espacios__icontains=q_sin_espacios) |
            Q(serie__autor__nombre__icontains=q) |
            Q(autor_sin_espacios__icontains=q_sin_espacios) |
            Q(editorial__nombre__icontains=q) |
            Q(editorial_sin_espacios__icontains=q_sin_espacios) |
            Q(serie__demografia__nombre__icontains=q) |
            Q(isbn__icontains=q) |
            Q(isbn_sin_guiones__icontains=q_sin_espacios)
        )

        # Si el usuario escribió varias palabras (ej: 'chainsaw man' o 'man chainsaw')
        palabras = q.split()
        if len(palabras) > 1:
            for p in palabras:
                p_clean = p.replace('-', '')
                filtro_texto |= (
                    Q(serie__titulo__icontains=p) |
                    Q(serie__autor__nombre__icontains=p) |
                    Q(editorial__nombre__icontains=p) |
                    Q(isbn__icontains=p_clean)
                )

        queryset = queryset.filter(filtro_texto).distinct()

    # Filtro por Serie
    serie_id = request.GET.get('serie')
    if serie_id:
        queryset = queryset.filter(serie_id=serie_id)

    # Filtro por Editorial
    editorial_id = request.GET.get('editorial')
    if editorial_id:
        queryset = queryset.filter(editorial_id=editorial_id)

    # Filtro por Demografía
    demografia_id = request.GET.get('demografia')
    if demografia_id:
        queryset = queryset.filter(serie__demografia_id=demografia_id)

    # Filtro por Stock
    stock_estado = request.GET.get('stock')
    if stock_estado == 'disponible':
        queryset = queryset.filter(stock__gt=0)
    elif stock_estado == 'agotado':
        queryset = queryset.filter(stock=0)

    return queryset

# 2. Leer / Listar: Muestra la tabla con filtros y paginación (Requisitos 11 y 13.a)
def listar_tomos(request):
    tomos_list = obtener_tomos_filtrados(request)
    
    # Paginación: 5 registros por página (Requisito 13.a)
    paginator = Paginator(tomos_list, 5)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    # Preservar filtros GET activos en la navegación de páginas
    query_dict = request.GET.copy()
    if 'page' in query_dict:
        del query_dict['page']
    query_string = query_dict.urlencode()
    
    # Opciones para los filtros desplegables
    series = Serie.objects.all().order_by('titulo')
    editoriales = Editorial.objects.all().order_by('nombre')
    demografias = Demografia.objects.all().order_by('nombre')
    
    context = {
        'page_obj': page_obj,
        'tomos': page_obj, # Para compatibilidad con templates
        'total_registros': tomos_list.count(),
        'query_string': query_string,
        'series': series,
        'editoriales': editoriales,
        'demografias': demografias,
        # Parámetros seleccionados para mantener estado en el formulario
        'filtro_q': request.GET.get('q', ''),
        'filtro_serie': request.GET.get('serie', ''),
        'filtro_editorial': request.GET.get('editorial', ''),
        'filtro_demografia': request.GET.get('demografia', ''),
        'filtro_stock': request.GET.get('stock', ''),
    }
    return render(request, 'tomos/listar.html', context)

# Exportar a Excel con formato profesional (Requisito 12)
def exportar_tomos_excel(request):
    tomos = obtener_tomos_filtrados(request)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario Manga"
    
    # Estilos profesionales
    header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D4D4D8'),
        right=Side(style='thin', color='D4D4D8'),
        top=Side(style='thin', color='D4D4D8'),
        bottom=Side(style='thin', color='D4D4D8')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    headers = [
        "ID", "Serie", "N° Tomo", "Editorial", "Demografía",
        "Autor", "ISBN", "Precio (CLP)", "Stock", "Fecha Registro"
    ]
    ws.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 26
    
    for row_idx, tomo in enumerate(tomos, start=2):
        row_data = [
            tomo.id,
            tomo.serie.titulo,
            f"Vol. #{tomo.numero_tomo}",
            tomo.editorial.nombre,
            tomo.serie.demografia.nombre if tomo.serie.demografia else "N/A",
            tomo.serie.autor.nombre if tomo.serie.autor else "N/A",
            tomo.isbn,
            int(tomo.precio),
            tomo.stock,
            tomo.fecha_ingreso.strftime("%d/%m/%Y") if tomo.fecha_ingreso else "N/A"
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 3, 7, 9, 10]:
                cell.alignment = center_align
            elif col_idx == 8:
                cell.alignment = right_align
                cell.number_format = '"$"#,##0'
            else:
                cell.alignment = left_align
                
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)
        
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="inventario_manga_tomos.xlsx"'
    wb.save(response)
    return response

# Exportar a CSV con UTF-8 BOM para soporte total en Excel (Requisito 13.b)
def exportar_tomos_csv(request):
    tomos = obtener_tomos_filtrados(request)
    
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="inventario_manga_tomos.csv"'
    
    writer = csv.writer(response, delimiter=';')
    writer.writerow([
        "ID", "Serie", "Numero_Tomo", "Editorial", "Demografia",
        "Autor", "ISBN", "Precio", "Stock", "Fecha_Registro"
    ])
    
    for tomo in tomos:
        writer.writerow([
            tomo.id,
            tomo.serie.titulo,
            tomo.numero_tomo,
            tomo.editorial.nombre,
            tomo.serie.demografia.nombre if tomo.serie.demografia else "",
            tomo.serie.autor.nombre if tomo.serie.autor else "",
            tomo.isbn,
            int(tomo.precio),
            tomo.stock,
            tomo.fecha_ingreso.strftime("%d/%m/%Y") if tomo.fecha_ingreso else ""
        ])
        
    return response

# 3. Crear Tomo: Procesa el formulario con el archivo adjunto
@solo_admin
def crear_tomo(request):
    if request.method == 'POST':
        # request.FILES es indispensable para recibir la imagen o archivo
        form = TomoForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Tomo agregado exitosamente al inventario.')
            return redirect('listar_tomos')
    else:
        form = TomoForm()
    return render(request, 'tomos/formulario.html', {'form': form, 'titulo': 'Agregar Nuevo Tomo'})

# 4. Actualizar / Editar Tomo
@solo_admin
def editar_tomo(request, pk):
    tomo = get_object_or_404(Tomo, pk=pk)
    if request.method == 'POST':
        form = TomoForm(request.POST, request.FILES, instance=tomo)
        if form.is_valid():
            form.save()
            messages.success(request, 'Tomo actualizado correctamente.')
            return redirect('listar_tomos')
    else:
        form = TomoForm(instance=tomo)
    return render(request, 'tomos/formulario.html', {'form': form, 'titulo': f'Editar Tomo #{tomo.numero_tomo} - {tomo.serie.titulo}'})

# 5. Eliminar Tomo
@solo_admin
def eliminar_tomo(request, pk):
    tomo = get_object_or_404(Tomo, pk=pk)
    if request.method == 'POST':
        tomo.delete()
        messages.success(request, 'Tomo eliminado del inventario.')
        return redirect('listar_tomos')
    return render(request, 'tomos/confirmar_eliminar.html', {'tomo': tomo})


# ==========================================
# CRUD: SERIE
# ==========================================
def listar_series(request):
    series = Serie.objects.select_related('demografia', 'autor').all()
    return render(request, 'series/listar.html', {'series': series})

@solo_admin
def crear_serie(request):
    if request.method == 'POST':
        form = SerieForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Serie creada exitosamente.')
            return redirect('listar_series')
    else:
        form = SerieForm()
    return render(request, 'series/formulario.html', {'form': form, 'titulo': 'Agregar Nueva Serie'})

@solo_admin
def editar_serie(request, pk):
    serie = get_object_or_404(Serie, pk=pk)
    if request.method == 'POST':
        form = SerieForm(request.POST, instance=serie)
        if form.is_valid():
            form.save()
            messages.success(request, 'Serie actualizada correctamente.')
            return redirect('listar_series')
    else:
        form = SerieForm(instance=serie)
    return render(request, 'series/formulario.html', {'form': form, 'titulo': f'Editar Serie: {serie.titulo}'})

@solo_admin
def eliminar_serie(request, pk):
    serie = get_object_or_404(Serie, pk=pk)
    if request.method == 'POST':
        serie.delete()
        messages.success(request, 'Serie eliminada correctamente.')
        return redirect('listar_series')
    return render(request, 'series/confirmar_eliminar.html', {'serie': serie})


# ==========================================
# CRUD: AUTOR
# ==========================================
def listar_autores(request):
    autores = Autor.objects.all()
    return render(request, 'autores/listar.html', {'autores': autores})

@solo_admin
def crear_autor(request):
    if request.method == 'POST':
        form = AutorForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Autor agregado exitosamente.')
            return redirect('listar_autores')
    else:
        form = AutorForm()
    return render(request, 'autores/formulario.html', {'form': form, 'titulo': 'Agregar Nuevo Autor'})

@solo_admin
def editar_autor(request, pk):
    autor = get_object_or_404(Autor, pk=pk)
    if request.method == 'POST':
        form = AutorForm(request.POST, instance=autor)
        if form.is_valid():
            form.save()
            messages.success(request, 'Autor actualizado correctamente.')
            return redirect('listar_autores')
    else:
        form = AutorForm(instance=autor)
    return render(request, 'autores/formulario.html', {'form': form, 'titulo': f'Editar Autor: {autor.nombre}'})

@solo_admin
def eliminar_autor(request, pk):
    autor = get_object_or_404(Autor, pk=pk)
    if request.method == 'POST':
        autor.delete()
        messages.success(request, 'Autor eliminado correctamente.')
        return redirect('listar_autores')
    return render(request, 'autores/confirmar_eliminar.html', {'autor': autor})


# ==========================================
# CRUD: EDITORIAL
# ==========================================
def listar_editoriales(request):
    editoriales = Editorial.objects.all()
    return render(request, 'editoriales/listar.html', {'editoriales': editoriales})

@solo_admin
def crear_editorial(request):
    if request.method == 'POST':
        form = EditorialForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Editorial agregada exitosamente.')
            return redirect('listar_editoriales')
    else:
        form = EditorialForm()
    return render(request, 'editoriales/formulario.html', {'form': form, 'titulo': 'Agregar Nueva Editorial'})

@solo_admin
def editar_editorial(request, pk):
    editorial = get_object_or_404(Editorial, pk=pk)
    if request.method == 'POST':
        form = EditorialForm(request.POST, instance=editorial)
        if form.is_valid():
            form.save()
            messages.success(request, 'Editorial actualizada correctamente.')
            return redirect('listar_editoriales')
    else:
        form = EditorialForm(instance=editorial)
    return render(request, 'editoriales/formulario.html', {'form': form, 'titulo': f'Editar Editorial: {editorial.nombre}'})

@solo_admin
def eliminar_editorial(request, pk):
    editorial = get_object_or_404(Editorial, pk=pk)
    if request.method == 'POST':
        editorial.delete()
        messages.success(request, 'Editorial eliminada correctamente.')
        return redirect('listar_editoriales')
    return render(request, 'editoriales/confirmar_eliminar.html', {'editorial': editorial})


# ==========================================
# CRUD: DEMOGRAFIA
# ==========================================
def listar_demografias(request):
    demografias = Demografia.objects.all()
    return render(request, 'demografias/listar.html', {'demografias': demografias})

@solo_admin
def crear_demografia(request):
    if request.method == 'POST':
        form = DemografiaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Demografía agregada exitosamente.')
            return redirect('listar_demografias')
    else:
        form = DemografiaForm()
    return render(request, 'demografias/formulario.html', {'form': form, 'titulo': 'Agregar Nueva Demografía'})

@solo_admin
def editar_demografia(request, pk):
    demografia = get_object_or_404(Demografia, pk=pk)
    if request.method == 'POST':
        form = DemografiaForm(request.POST, instance=demografia)
        if form.is_valid():
            form.save()
            messages.success(request, 'Demografía actualizada correctamente.')
            return redirect('listar_demografias')
    else:
        form = DemografiaForm(instance=demografia)
    return render(request, 'demografias/formulario.html', {'form': form, 'titulo': f'Editar Demografía: {demografia.nombre}'})

@solo_admin
def eliminar_demografia(request, pk):
    demografia = get_object_or_404(Demografia, pk=pk)
    if request.method == 'POST':
        demografia.delete()
        messages.success(request, 'Demografía eliminada correctamente.')
        return redirect('listar_demografias')
    return render(request, 'demografias/confirmar_eliminar.html', {'demografia': demografia})


# ==========================================
# AUTENTICACIÓN Y CONTROL DE ACCESO
# ==========================================
def iniciar_sesion(request):
    if request.user.is_authenticated:
        return redirect('inicio')
    
    if request.method == 'POST':
        form = LoginForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            auth_login(request, user)
            rol = "Administrador" if (user.is_staff or user.is_superuser) else "Operador"
            messages.success(request, f'Sesión iniciada como {user.username} ({rol}).')
            next_url = request.GET.get('next') or 'inicio'
            return redirect(next_url)
        else:
            messages.error(request, 'Usuario o contraseña incorrectos. Por favor intenta de nuevo.')
    else:
        form = LoginForm()
    
    return render(request, 'auth/login.html', {'form': form})

def cerrar_sesion(request):
    if request.user.is_authenticated:
        auth_logout(request)
        messages.info(request, 'Has cerrado sesión exitosamente.')
    return redirect('inicio')

def registro_usuario(request):
    if request.user.is_authenticated:
        return redirect('inicio')
    
    if request.method == 'POST':
        form = RegistroUsuarioForm(request.POST)
        if form.is_valid():
            nuevo_usuario = form.save(commit=False)
            # Todo nuevo usuario registrado inicia como Operador (sin permisos de staff/admin)
            nuevo_usuario.is_staff = False
            nuevo_usuario.is_superuser = False
            nuevo_usuario.save()
            messages.success(
                request,
                f'¡Cuenta creada con éxito para {nuevo_usuario.username}! Ya puedes iniciar sesión con tu usuario.'
            )
            return redirect('login')
    else:
        form = RegistroUsuarioForm()
    
    return render(request, 'auth/registro.html', {'form': form})
